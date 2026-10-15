"""Build the narrated presentation from the deck's speaker notes (Windows + Microsoft PowerPoint).

    python scripts/build_narrated_deck.py [--voice zira|david] [--rate -1]

Steps: read each slide's speaker notes, synthesize speech with the built-in Windows voices (SAPI), embed one
audio clip per slide with automatic advance, then export Topic_7_Capstone_Presentation_Narrated.pptx and .mp4.
The notes are the narration script (see narration_script.md), so the same text can be recorded in your own voice
with PowerPoint's Record Slide Show instead.
"""
import argparse
import os
import re
import time
import wave

import win32com.client
from pptx import Presentation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "Topic_7_Capstone_Presentation.pptx")
OUT_PPTX = os.path.join(ROOT, "Topic_7_Capstone_Presentation_Narrated.pptx")
OUT_MP4 = os.path.join(ROOT, "Topic_7_Capstone_Presentation_Narrated.mp4")
AUDIO_DIR = os.path.join(ROOT, "artifacts", "narration")
SCRIPT_MD = os.path.join(ROOT, "narration_script.md")

# Spoken forms for terms the Windows voices mispronounce (applied to the speech text only, not the notes).
SPOKEN = [
    (r"\bABPM\b", "A B P M"), (r"\bTreeSHAP\b", "Tree Shap"), (r"\bSHAP\b", "Shap"), (r"\bXGBoost\b", "X G Boost"),
    (r"\bROC-AUC\b", "R O C A U C"), (r"\bAUC\b", "A U C"), (r"\bECE\b", "E C E"), (r"\bJWT\b", "J W T"),
    (r"\bFastAPI\b", "Fast A P I"), (r"\bAPI\b", "A P I"), (r"\bPostgreSQL\b", "Postgres Q L"), (r"\bTRIPOD\+AI\b", "TRIPOD plus A I"),
    (r"\bHIPAA\b", "HIPAA"), (r"topic7-submission", "topic 7 submission"), (r"\bz about 12\b", "z-score about 12"),
    (r"\bz -3\.1\b", "z-score minus 3 point 1"), (r"\(-(\d+\.\d+) to \+(\d+\.\d+)\)", r"(minus \1 to plus \2)"),
    (r"\bmmHg\b", "millimetres of mercury"), (r"\b(\d+)ms\b", r"\1 milliseconds"),
]


def spoken(text):
    for pat, rep in SPOKEN:
        text = re.sub(pat, rep, text)
    return text


def synthesize(text, path, voice_index, rate):
    voice = win32com.client.Dispatch("SAPI.SpVoice")
    voice.Voice = voice.GetVoices().Item(voice_index)
    voice.Rate = rate
    fmt = win32com.client.Dispatch("SAPI.SpAudioFormat"); fmt.Type = 22          # 22 kHz, 16-bit, mono
    stream = win32com.client.Dispatch("SAPI.SpFileStream"); stream.Format = fmt
    stream.Open(path, 3)                                                           # create for write
    voice.AudioOutputStream = stream
    voice.Speak(text, 0)
    stream.Close()
    with wave.open(path) as w:
        return w.getnframes() / w.getframerate()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--voice", default="zira", choices=["zira", "david"])
    ap.add_argument("--rate", type=int, default=-1)
    args = ap.parse_args()
    vidx = {"david": 0, "zira": 1}[args.voice]

    os.makedirs(AUDIO_DIR, exist_ok=True)
    notes = [s.notes_slide.notes_text_frame.text.strip() for s in Presentation(SRC).slides]
    titles = []
    for s in Presentation(SRC).slides:
        t = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        titles.append(t[0].split("\n")[0] if t else "")

    with open(SCRIPT_MD, "w", encoding="utf-8") as f:
        f.write("# Narration script\n\nOne paragraph per slide (the deck's speaker notes). Read aloud, about 6 to 7 minutes in total.\n\n")
        for i, (t, n) in enumerate(zip(titles, notes), 1):
            f.write(f"## Slide {i}\n\n{n}\n\n")

    durations = []
    for i, n in enumerate(notes, 1):
        path = os.path.join(AUDIO_DIR, f"slide{i:02d}.wav")
        d = synthesize(spoken(n), path, vidx, args.rate)
        durations.append(d)
        print(f"slide {i:2d}: {d:5.1f} s  ({len(n.split())} words)")
    print(f"total narration: {sum(durations) / 60:.1f} min")

    app = win32com.client.Dispatch("PowerPoint.Application")
    try:
        pres = app.Presentations.Open(SRC, WithWindow=False)
        for i, d in enumerate(durations, 1):
            slide = pres.Slides(i)
            snd = slide.Shapes.AddMediaObject2(os.path.join(AUDIO_DIR, f"slide{i:02d}.wav"), False, True, -300, -300, 40, 40)   # off-slide icon
            snd.Name = f"Narration {i}"
            snd.AnimationSettings.PlaySettings.PlayOnEntry = True
            snd.AnimationSettings.PlaySettings.HideWhileNotPlaying = False
            slide.SlideShowTransition.AdvanceOnTime = True
            slide.SlideShowTransition.AdvanceTime = d + 1.5
        if os.path.exists(OUT_PPTX):
            os.remove(OUT_PPTX)
        pres.SaveAs(OUT_PPTX)
        if os.path.exists(OUT_MP4):
            os.remove(OUT_MP4)
        pres.CreateVideo(OUT_MP4, True, 5, 720, 24, 40)
        t0 = time.time()
        while pres.CreateVideoStatus in (1, 2):                                   # in progress / queued
            time.sleep(3)
            if time.time() - t0 > 1500:
                raise TimeoutError("video export timed out")
        print("video status:", pres.CreateVideoStatus, "(3 = done)", f"in {time.time() - t0:.0f}s")
        pres.Close()
    finally:
        try:
            app.Quit()
        except Exception:
            pass
    for p in (OUT_PPTX, OUT_MP4):
        print(p, f"{os.path.getsize(p) / 1e6:.1f} MB" if os.path.exists(p) else "MISSING")


if __name__ == "__main__":
    main()
